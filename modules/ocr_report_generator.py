from pathlib import Path
from datetime import datetime
import json

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def load_json(path):
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"Required file not found: {file_path}")
    with open(file_path, "r", encoding="utf-8") as f:
        return json.load(f)


def setup_document_style(document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        if style_name in styles:
            style = styles[style_name]
            style.font.name = "Microsoft YaHei"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_center_title(document, title):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(20)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_summary_table(document, rows):
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Value"

    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)

    document.add_paragraph("")


def add_detail_table(document, ocr_results):
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"

    headers = [
        "No.",
        "File Name",
        "OCR",
        "Language",
        "Raw Length",
        "Cleaned Length",
        "Error"
    ]

    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header

    for index, item in enumerate(ocr_results, start=1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = str(item.get("file_name", ""))
        cells[2].text = "Success" if item.get("ocr_success") else "Failed"
        cells[3].text = str(item.get("ocr_language", ""))
        cells[4].text = str(item.get("text_length", ""))
        cells[5].text = str(item.get("cleaned_text_length", ""))
        cells[6].text = str(item.get("error_message", ""))[:80]

    document.add_paragraph("")


def find_basic_info(file_name, basic_results):
    for item in basic_results:
        if item.get("file_name") == file_name:
            return item
    return {}


def add_text_block(document, title, text):
    document.add_heading(title, level=3)

    if not text:
        document.add_paragraph("No text extracted.")
        return

    paragraphs = text.split("\n\n")

    for paragraph in paragraphs:
        paragraph = paragraph.strip()
        if paragraph:
            document.add_paragraph(paragraph)


def generate_ocr_word_report(
    basic_json="output/image_basic_info.json",
    cleaned_ocr_json="output/image_ocr_results_cleaned.json",
    output_docx="output/image_ocr_report.docx"
):
    basic_results = load_json(basic_json)
    ocr_results = load_json(cleaned_ocr_json)

    output_path = Path(output_docx)
    output_path.parent.mkdir(exist_ok=True)

    total_files = len(ocr_results)
    success_count = sum(1 for item in ocr_results if item.get("ocr_success"))
    failed_count = total_files - success_count
    text_found_count = sum(1 for item in ocr_results if item.get("cleaned_text_length", 0) > 0)

    document = Document()
    setup_document_style(document)

    add_center_title(document, "Batch Media Insight Extractor")
    document.add_paragraph("OCR Comprehensive Report")
    document.add_paragraph(
        "This report combines image preview, basic image metadata, OCR extraction status, "
        "and cleaned OCR text for easier reading and later analysis."
    )

    document.add_heading("1. Report Summary", level=1)

    summary_rows = [
        ("Generated Time", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total Images", total_files),
        ("OCR Success", success_count),
        ("OCR Failed", failed_count),
        ("Images With Text", text_found_count),
        ("Basic Info Source", basic_json),
        ("Cleaned OCR Source", cleaned_ocr_json),
    ]

    add_summary_table(document, summary_rows)

    document.add_heading("2. OCR Overview Table", level=1)
    add_detail_table(document, ocr_results)

    document.add_heading("3. Image OCR Details", level=1)

    for index, item in enumerate(ocr_results, start=1):
        file_name = item.get("file_name", "")
        basic = find_basic_info(file_name, basic_results)

        document.add_heading(f"{index}. {file_name}", level=2)

        info_rows = [
            ("File Name", file_name),
            ("Format", basic.get("image_format", "")),
            ("Dimension", f"{basic.get('width', '')} x {basic.get('height', '')}"),
            ("Size MB", basic.get("size_mb", "")),
            ("OCR Success", item.get("ocr_success", "")),
            ("OCR Language", item.get("ocr_language", "")),
            ("Raw Text Length", item.get("text_length", "")),
            ("Cleaned Text Length", item.get("cleaned_text_length", "")),
        ]

        add_summary_table(document, info_rows)

        full_path = basic.get("full_path") or item.get("full_path", "")

        try:
            image_path = Path(full_path)
            if image_path.exists():
                document.add_paragraph("Image Preview:")
                document.add_picture(str(image_path), width=Inches(2.6))
            else:
                document.add_paragraph("Image preview unavailable: file path does not exist.")
        except Exception as e:
            document.add_paragraph(f"Image preview unavailable: {e}")

        document.add_paragraph("")

        if item.get("error_message"):
            document.add_paragraph(f"OCR Error: {item.get('error_message')}")

        add_text_block(
            document,
            "Cleaned OCR Text",
            item.get("cleaned_text", "")
        )

        document.add_page_break()

    document.add_heading("4. Next Development Steps", level=1)
    document.add_paragraph("Next module: AI-based OCR text summarization and keyword extraction.")
    document.add_paragraph("Future module: video audio extraction and speech-to-text transcription.")
    document.add_paragraph("Future module: local web software interface for easier batch processing.")

    document.save(output_path)

    return output_path
