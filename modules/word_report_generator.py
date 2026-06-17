from pathlib import Path
from datetime import datetime
import json

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def load_image_info(json_path: str = "output/image_basic_info.json"):
    path = Path(json_path)

    if not path.exists():
        raise FileNotFoundError(
            "image_basic_info.json not found. Please run run_step_002_image_info.py first."
        )

    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def add_title(document: Document, title: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(20)


def add_key_value_table(document: Document, rows):
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    header_cells[0].text = "Item"
    header_cells[1].text = "Value"

    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)

    document.add_paragraph("")


def add_image_detail_table(document: Document, results):
    table = document.add_table(rows=1, cols=8)
    table.style = "Table Grid"

    headers = [
        "No.",
        "File Name",
        "Format",
        "Size MB",
        "Width",
        "Height",
        "Readable",
        "EXIF Time"
    ]

    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header

    for i, item in enumerate(results, start=1):
        cells = table.add_row().cells
        cells[0].text = str(i)
        cells[1].text = str(item.get("file_name", ""))
        cells[2].text = str(item.get("image_format", ""))
        cells[3].text = str(item.get("size_mb", ""))
        cells[4].text = str(item.get("width", ""))
        cells[5].text = str(item.get("height", ""))
        cells[6].text = "Yes" if item.get("is_readable") else "No"
        cells[7].text = str(item.get("exif_datetime", ""))

    document.add_paragraph("")


def add_image_preview_section(document: Document, results):
    document.add_heading("Image Preview", level=1)

    for i, item in enumerate(results, start=1):
        file_name = item.get("file_name", "")
        full_path = item.get("full_path", "")
        readable = item.get("is_readable", False)

        document.add_heading(f"{i}. {file_name}", level=2)

        info_text = (
            f"Format: {item.get('image_format', '')} | "
            f"Size: {item.get('size_mb', '')} MB | "
            f"Dimension: {item.get('width', '')} x {item.get('height', '')}"
        )
        document.add_paragraph(info_text)

        if not readable:
            document.add_paragraph(f"Unreadable image. Error: {item.get('error_message', '')}")
            continue

        try:
            image_path = Path(full_path)
            if image_path.exists():
                document.add_picture(str(image_path), width=Inches(2.2))
            else:
                document.add_paragraph("Preview unavailable: file path does not exist.")
        except Exception as e:
            document.add_paragraph(f"Preview unavailable: {e}")

        document.add_paragraph("")


def generate_image_word_report(
    input_json: str = "output/image_basic_info.json",
    output_docx: str = "output/image_basic_report.docx"
):
    results = load_image_info(input_json)

    total_images = len(results)
    readable_images = sum(1 for item in results if item.get("is_readable"))
    unreadable_images = total_images - readable_images

    output_path = Path(output_docx)
    output_path.parent.mkdir(exist_ok=True)

    document = Document()

    add_title(document, "Batch Media Insight Extractor")
    document.add_paragraph("Image Basic Information Report")

    document.add_paragraph(
        "This report summarizes the basic information extracted from batch image files, "
        "including file format, file size, image dimensions, readability, and basic EXIF fields."
    )

    document.add_heading("1. Report Summary", level=1)

    summary_rows = [
        ("Generated Time", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total Images", total_images),
        ("Readable Images", readable_images),
        ("Unreadable Images", unreadable_images),
        ("Input Source", "input_media"),
        ("Data Source", input_json),
    ]

    add_key_value_table(document, summary_rows)

    document.add_heading("2. Image Detail Table", level=1)

    if total_images == 0:
        document.add_paragraph("No image files were found.")
    else:
        add_image_detail_table(document, results)

    document.add_heading("3. Error Records", level=1)

    error_items = [item for item in results if not item.get("is_readable")]

    if not error_items:
        document.add_paragraph("No unreadable or damaged image files were detected.")
    else:
        for item in error_items:
            document.add_paragraph(
                f"{item.get('file_name', '')}: {item.get('error_message', '')}"
            )

    if total_images > 0:
        add_image_preview_section(document, results)

    document.add_heading("4. Next Development Steps", level=1)
    document.add_paragraph("Next module: OCR text extraction from images.")
    document.add_paragraph("Future module: video audio extraction and speech-to-text transcription.")
    document.add_paragraph("Future module: PDF export and local software interface.")

    document.save(output_path)

    return output_path


if __name__ == "__main__":
    path = generate_image_word_report()
    print(f"Word report generated: {path}")
