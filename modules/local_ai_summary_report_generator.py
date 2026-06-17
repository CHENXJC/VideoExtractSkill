from pathlib import Path
from datetime import datetime
import json

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def load_json(path):
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"Required file not found: {file_path}")

    with open(file_path, "r", encoding="utf-8") as f:
        return json.load(f)


def setup_document_style(document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        if style_name in styles:
            style = styles[style_name]
            style.font.name = "Microsoft YaHei"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_center_title(document, title):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(20)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_key_value_table(document, rows):
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Value"

    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)

    document.add_paragraph("")


def add_overview_table(document, results):
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"

    headers = [
        "No.",
        "File Name",
        "Content Type",
        "Keywords",
        "Text Length",
        "Type Score"
    ]

    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header

    for index, item in enumerate(results, start=1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = str(item.get("file_name", ""))
        cells[2].text = str(item.get("content_type", ""))
        cells[3].text = "、".join(item.get("keywords", [])[:5])
        cells[4].text = str(item.get("cleaned_text_length", ""))
        cells[5].text = str(item.get("content_type_score", ""))

    document.add_paragraph("")


def find_basic_info(file_name, basic_results):
    for item in basic_results:
        if item.get("file_name") == file_name:
            return item
    return {}


def add_bullet_points(document, points):
    if not points:
        document.add_paragraph("No key points extracted.")
        return

    for index, point in enumerate(points, start=1):
        document.add_paragraph(f"{index}. {point}")


def generate_local_summary_word_report(
    summary_json="output/image_ai_summary_local.json",
    basic_json="output/image_basic_info.json",
    output_docx="output/reports/image_ai_summary_local_report.docx"
):
    summary_results = load_json(summary_json)
    basic_results = load_json(basic_json)

    output_path = Path(output_docx)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    total_files = len(summary_results)
    content_types = {}

    for item in summary_results:
        content_type = item.get("content_type", "Unknown")
        content_types[content_type] = content_types.get(content_type, 0) + 1

    document = Document()
    setup_document_style(document)

    add_center_title(document, "Batch Media Insight Extractor")
    document.add_paragraph("Local AI Summary Report")
    document.add_paragraph(
        "This report summarizes OCR-extracted image text using a local rule-based analysis module. "
        "It includes content type detection, keyword extraction, summary generation, key points, and archive suggestions."
    )

    document.add_heading("1. Report Summary", level=1)

    type_summary = "; ".join([f"{key}: {value}" for key, value in content_types.items()])

    summary_rows = [
        ("Generated Time", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total Files", total_files),
        ("Analysis Mode", "Local rule-based AI Summary"),
        ("Content Type Distribution", type_summary),
        ("Summary Source", summary_json),
        ("Basic Info Source", basic_json),
    ]

    add_key_value_table(document, summary_rows)

    document.add_heading("2. AI Summary Overview", level=1)
    add_overview_table(document, summary_results)

    document.add_heading("3. File-Level Analysis", level=1)

    for index, item in enumerate(summary_results, start=1):
        file_name = item.get("file_name", "")
        basic = find_basic_info(file_name, basic_results)

        document.add_heading(f"{index}. {file_name}", level=2)

        info_rows = [
            ("File Name", file_name),
            ("Content Type", item.get("content_type", "")),
            ("Content Type Score", item.get("content_type_score", "")),
            ("OCR Language", item.get("ocr_language", "")),
            ("Cleaned Text Length", item.get("cleaned_text_length", "")),
            ("Keywords", "、".join(item.get("keywords", []))),
        ]

        add_key_value_table(document, info_rows)

        full_path = basic.get("full_path", "")

        try:
            image_path = Path(full_path)
            if image_path.exists():
                document.add_paragraph("Image Preview:")
                document.add_picture(str(image_path), width=Inches(2.5))
            else:
                document.add_paragraph("Image preview unavailable: file path does not exist.")
        except Exception as e:
            document.add_paragraph(f"Image preview unavailable: {e}")

        document.add_heading("Summary", level=3)
        document.add_paragraph(item.get("summary", "") or "No summary generated.")

        document.add_heading("Key Points", level=3)
        add_bullet_points(document, item.get("key_points", []))

        document.add_heading("Archive Suggestion", level=3)
        document.add_paragraph(item.get("archive_suggestion", "") or "No archive suggestion generated.")

        document.add_page_break()

    document.add_heading("4. Next Development Steps", level=1)
    document.add_paragraph("Next module: OpenAI / ChatGPT enhanced summarization.")
    document.add_paragraph("Future module: video audio extraction and video OCR.")
    document.add_paragraph("Future module: local web software interface for batch upload and one-click export.")

    document.save(output_path)

    return output_path
