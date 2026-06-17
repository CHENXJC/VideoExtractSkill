from pathlib import Path
from datetime import datetime
import json
import win32com.client

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def load_json(path):
    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(f"Required file not found: {file_path}")

    with open(file_path, "r", encoding="utf-8") as f:
        return json.load(f)


def setup_document_style(document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        if style_name in styles:
            style = styles[style_name]
            style.font.name = "Microsoft YaHei"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_title(document, title):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(20)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_key_value_table(document, rows):
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Value"

    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)

    document.add_paragraph("")


def add_overview_table(document, results):
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"

    headers = ["No.", "File Name", "Language", "Content Type", "Text Length", "Type Score"]

    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header

    for i, item in enumerate(results, start=1):
        cells = table.add_row().cells
        cells[0].text = str(i)
        cells[1].text = str(item.get("file_name", ""))
        cells[2].text = str(item.get("detected_language", ""))
        cells[3].text = str(item.get("content_type", ""))
        cells[4].text = str(item.get("cleaned_text_length", ""))
        cells[5].text = str(item.get("content_type_score", ""))

    document.add_paragraph("")


def add_bullets(document, points):
    if not points:
        document.add_paragraph("No key points extracted.")
        return

    for index, point in enumerate(points, start=1):
        document.add_paragraph(f"{index}. {point}")


def convert_docx_to_pdf(docx_path, pdf_path):
    docx_file = Path(docx_path).resolve()
    pdf_file = Path(pdf_path).resolve()

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False

    try:
        doc = word.Documents.Open(str(docx_file))
        doc.SaveAs(str(pdf_file), FileFormat=17)
        doc.Close()
    finally:
        word.Quit()

    return pdf_file


def generate_video_summary_report(
    summary_json="output/video_summary_local.json",
    output_docx="output/reports/video_summary_local_report.docx",
    output_pdf="output/reports/video_summary_local_report.pdf",
    report_format="both"
):
    results = load_json(summary_json)

    reports_dir = Path(output_docx).parent
    reports_dir.mkdir(parents=True, exist_ok=True)

    total_files = len(results)
    success_count = sum(1 for item in results if item.get("transcription_success"))
    text_count = sum(1 for item in results if item.get("cleaned_text_length", 0) > 0)

    document = Document()
    setup_document_style(document)

    add_title(document, "Batch Media Insight Extractor")
    document.add_paragraph("Video Transcript Summary Report")
    document.add_paragraph(
        "This report summarizes video speech-to-text transcripts using a local rule-based summary module."
    )

    document.add_heading("1. Report Summary", level=1)

    add_key_value_table(document, [
        ("Generated Time", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total Videos", total_files),
        ("Transcription Success", success_count),
        ("Videos With Text", text_count),
        ("Analysis Mode", "Local rule-based video transcript summary"),
        ("Data Source", summary_json),
    ])

    document.add_heading("2. Overview", level=1)
    add_overview_table(document, results)

    document.add_heading("3. Video-Level Analysis", level=1)

    for index, item in enumerate(results, start=1):
        document.add_heading(f"{index}. {item.get('file_name', '')}", level=2)

        add_key_value_table(document, [
            ("File Name", item.get("file_name", "")),
            ("Language", item.get("detected_language", "")),
            ("Language Probability", item.get("language_probability", "")),
            ("Content Type", item.get("content_type", "")),
            ("Content Type Score", item.get("content_type_score", "")),
            ("Raw Text Length", item.get("raw_text_length", "")),
            ("Cleaned Text Length", item.get("cleaned_text_length", "")),
            ("Keywords", "、".join(item.get("keywords", []))),
        ])

        document.add_heading("Summary", level=3)
        document.add_paragraph(item.get("summary", "") or "No summary generated.")

        document.add_heading("Key Points", level=3)
        add_bullets(document, item.get("key_points", []))

        document.add_heading("Action Suggestion", level=3)
        document.add_paragraph(item.get("action_suggestion", "") or "No suggestion generated.")

        document.add_heading("Cleaned Transcript", level=3)
        transcript = item.get("cleaned_text", "")
        if transcript:
            # 避免 Word 段落过长，按长度拆开
            chunk_size = 800
            for start in range(0, len(transcript), chunk_size):
                document.add_paragraph(transcript[start:start + chunk_size])
        else:
            document.add_paragraph("No transcript text found.")

        document.add_page_break()

    document.add_heading("4. Next Development Steps", level=1)
    document.add_paragraph("Next module: add video transcript summary into the Streamlit web app.")
    document.add_paragraph("Future module: OpenAI / ChatGPT enhanced transcript analysis.")
    document.add_paragraph("Future module: video frame OCR and combined video report.")

    document.save(output_docx)

    generated_docx = Path(output_docx)
    generated_pdf = None

    if report_format in ["pdf", "both"]:
        generated_pdf = convert_docx_to_pdf(output_docx, output_pdf)

    if report_format == "pdf":
        try:
            generated_docx.unlink()
            generated_docx = None
        except Exception:
            pass

    return generated_docx, generated_pdf
